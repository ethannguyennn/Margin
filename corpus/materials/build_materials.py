"""Build the four synthetic demo materials (Part 4).

Everything here is fictional training content for the Halcyon HX-2 portable
fluid pump, in the style of corpus/SOP-114-C.md. Never replace these with a
real employer document.

    python corpus/materials/build_materials.py

writes, next to this file:

    SOP-150-pump-seal-swap.pdf      text-layer PDF, 2 pages, warning box, exploded-view figure
    SOP-151-filter-change.docx      DOCX, torque table, embedded figure
    SOP-152-valve-lineup-photo.jpg  phone photo of a printed page: skew, shadow, noise
    SOP-153-quick-start.md          plain text, star pattern described in words, no figure

and the figure sources under figures/ (also used by the renderer's golden test).
Needs Pillow and python-docx. The PDF is written by hand (standard Helvetica,
text drawn with Tj), so pypdf reads a real text layer and no PDF library is needed.
"""

from __future__ import annotations

import io
import random
import zlib
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont

HERE = Path(__file__).resolve().parent
FONT = HERE.parents[1] / "relay" / "src" / "specs_relay" / "diagrams" / "fonts" / "AtkinsonHyperlegible-Bold.ttf"
FIG = HERE / "figures"


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT), size)


# --- figures (black line art on white paper) ------------------------------------------

def seal_housing_figure() -> Image.Image:
    """SOP-150 Figure 1: exploded side view (left) and gland nut order (right)."""
    im = Image.new("L", (900, 420), 255)
    d = ImageDraw.Draw(im)
    f, small = font(22), font(18)
    # Left: exploded side view, parts spread along the shaft axis.
    d.rectangle((40, 150, 170, 290), outline=0, width=4)          # seal housing
    d.line((170, 220, 460, 220), fill=0, width=2)                  # shaft axis
    d.ellipse((215, 170, 245, 270), outline=0, width=4)            # O-ring
    d.rectangle((290, 160, 315, 280), outline=0, width=4)          # gland plate
    for y in (175, 265):
        d.rectangle((360, y - 10, 390, y + 10), outline=0, width=4)  # gland nuts
    for x, y, n in ((105, 120, "4"), (230, 140, "3"), (302, 130, "2"), (375, 135, "1")):
        d.line((x, y + 14, x, y + 34), fill=0, width=2)
        d.ellipse((x - 14, y - 14, x + 14, y + 14), outline=0, width=3)
        d.text((x, y), n, font=small, fill=0, anchor="mm")
    legend = ["1  Gland nut M6 (4x)", "2  Gland plate", "3  O-ring HX2-311", "4  Seal housing"]
    for i, line in enumerate(legend):
        d.text((40, 320 + i * 24), line, font=small, fill=0)
    # Right: front view of the gland plate, nut tightening order.
    d.rectangle((560, 70, 840, 350), outline=0, width=4)
    d.ellipse((665, 175, 735, 245), outline=0, width=4)           # shaft bore
    order = {"1": (600, 110), "2": (800, 310), "3": (800, 110), "4": (600, 310)}
    for n, (x, y) in order.items():
        d.ellipse((x - 22, y - 22, x + 22, y + 22), outline=0, width=4)
        d.text((x, y), n, font=f, fill=0, anchor="mm")
    d.text((700, 375), "Tightening order", font=small, fill=0, anchor="mm")
    return im


def filter_figure() -> Image.Image:
    """SOP-151 Figure 1: filter bowl, element and the two clamp screws."""
    im = Image.new("L", (640, 400), 255)
    d = ImageDraw.Draw(im)
    small = font(18)
    d.rectangle((220, 40, 420, 90), outline=0, width=4)            # filter head
    d.rectangle((240, 90, 400, 330), outline=0, width=4)           # bowl
    d.rectangle((280, 110, 360, 300), outline=0, width=3)          # element
    for y in range(125, 300, 22):
        d.line((285, y, 355, y), fill=0, width=1)                  # pleats
    for x in (200, 440):
        d.rectangle((x - 12, 55, x + 12, 75), outline=0, width=3)  # clamp screws
    d.rectangle((305, 330, 335, 355), outline=0, width=3)          # drain plug
    labels = [((200, 30), "Clamp screw M5"), ((470, 200), "Element HX2-420"), ((320, 380), "Drain plug M10")]
    for (x, y), text in labels:
        d.text((x, y), text, font=small, fill=0, anchor="mm")
    d.line((440, 200, 360, 200), fill=0, width=2)
    return im


def valve_figure() -> Image.Image:
    """SOP-152 Figure 1: valve handle positions for flushing."""
    im = Image.new("L", (700, 260), 255)
    d = ImageDraw.Draw(im)
    f = font(20)
    d.line((30, 130, 670, 130), fill=0, width=6)                   # main pipe
    valves = [(110, "V1", "closed", 90), (270, "V3", "open", 0), (430, "V4", "45", 45), (590, "V2", "open", 0)]
    for x, name, state, angle in valves:
        d.polygon([(x - 26, 108), (x + 26, 152), (x + 26, 108), (x - 26, 152)], outline=0, width=3)
        dx, dy = {90: (0, 44), 0: (44, 0), 45: (31, 31)}[angle]
        d.line((x - dx, 130 - dy, x + dx, 130 + dy), fill=0, width=7)  # handle
        d.text((x, 40), name, font=f, fill=0, anchor="mm")
        d.text((x, 215), state, font=f, fill=0, anchor="mm")
    return im


# --- SOP-150: hand-written text-layer PDF ----------------------------------------------

def _pdf_escape(text: str) -> bytes:
    out = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return out.encode("cp1252")  # WinAnsiEncoding


def _wrap(text: str, width: int) -> list[str]:
    lines, line = [], ""
    for word in text.split():
        if line and len(line) + 1 + len(word) > width:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    return lines + ([line] if line else [])


class _Page:
    def __init__(self) -> None:
        self.ops: list[bytes] = []
        self.y = 742
        self.images: dict[str, tuple[bytes, int, int]] = {}

    def text(self, s: str, size: int = 10, bold: bool = False, indent: int = 0, wrap: int = 92) -> None:
        for line in _wrap(s, wrap):
            self.ops.append(b"BT /%s %d Tf %d %d Td (" % (b"F2" if bold else b"F1", size, 54 + indent, self.y)
                            + _pdf_escape(line) + b") Tj ET")
            self.y -= size + 4
        self.y -= 4

    def box(self, lines: list[str]) -> None:
        top = self.y + 14
        for s in lines:
            self.text(s, size=10, bold=True, indent=10, wrap=88)
        self.ops.append(b"2 w %d %d %d %d re S" % (54, self.y + 4, 504, top - self.y - 4))
        self.y -= 10

    def image(self, name: str, im: Image.Image, width_pt: int) -> None:
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=92)
        h_pt = round(width_pt * im.height / im.width)
        self.images[name] = (buf.getvalue(), im.width, im.height)
        self.y -= h_pt
        self.ops.append(b"q %d 0 0 %d %d %d cm /%s Do Q" % (width_pt, h_pt, 54, self.y, name.encode()))
        self.y -= 10


def write_pdf(pages: list[_Page]) -> bytes:
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    f1 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    f2 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    pages_id = len(objs) + 1 + 2 * len(pages) + sum(len(p.images) for p in pages)
    kids = []
    for p in pages:
        xobjs = []
        for name, (jpeg, w, h) in p.images.items():
            img = add(b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceGray "
                      b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % (w, h, len(jpeg))
                      + jpeg + b"\nendstream")
            xobjs.append(b"/%s %d 0 R" % (name.encode(), img))
        stream = zlib.compress(b"\n".join(p.ops))
        content = add(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream) + stream + b"\nendstream")
        resources = b"<< /Font << /F1 %d 0 R /F2 %d 0 R >> /XObject << %s >> >>" % (f1, f2, b" ".join(xobjs))
        kids.append(add(b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Resources %s /Contents %d 0 R >>"
                        % (pages_id, resources, content)))
    assert add(b"<< /Type /Pages /Kids [%s] /Count %d >>"
               % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))) == pages_id
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, catalog, xref)
    return bytes(out)


def build_sop150() -> bytes:
    p1, p2 = _Page(), _Page()
    p1.text("SOP-150 Pump Seal Swap", size=16, bold=True)
    p1.text("Halcyon HX-2 portable fluid pump. Revision A. Effective 2026-10-01. SYNTHETIC TRAINING "
            "DOCUMENT: all names, parts, limits and records are fictional demo data.", size=9)
    p1.text("1 Scope", size=12, bold=True)
    p1.text("Replace the shaft seal O-ring on an HX-2 pump that failed the leak check under SOP-140. "
            "No other seal or pump model is covered.")
    p1.text("2 Tools and parts", size=12, bold=True)
    p1.text("10 mm socket. Torque wrench, 1 N-m to 10 N-m range, checked under SOP-131. 0.20 mm feeler "
            "gauge. New O-ring HX2-311.")
    p1.box(["WARNING: Isolate power and drain the pump before opening the seal housing.",
            "Do not reuse the old O-ring. Do not lubricate the new O-ring."])
    p1.image("Im1", seal_housing_figure(), 432)
    p1.text("Figure 1. Seal housing: exploded view (left) and gland nut tightening order (right).", size=9)
    p1.text("3 Procedure", size=12, bold=True)
    p1.text("3.1 Remove the four M6 gland nuts with the 10 mm socket. Keep the nuts for reuse.")
    p1.text("3.2 Lift the gland plate off the shaft and discard the old O-ring.")
    p1.text("3.3 Fit the new O-ring HX2-311 into the groove of the seal housing. Fit it dry.")
    p2.text("3.4 Refit the gland plate and torque the four M6 gland nuts to 4 N-m in the order shown "
            "in Figure 1: upper left, lower right, upper right, lower left.")
    p2.text("3.5 Check the gap under the gland plate with the 0.20 mm feeler gauge at all four corners. "
            "The gauge must not enter.")
    p2.text("3.6 Restore power and run the pump for 30 s at 1200 rpm. Check the seal for drips. Any drip "
            "is a nonconformance: report it under SOP-140.")
    p2.text("4 Records", size=12, bold=True)
    p2.text("Form SW-150 is retained for 12 months. This procedure contains no pressure test.")
    return write_pdf([p1, p2])


# --- SOP-151: DOCX ---------------------------------------------------------------------

def build_sop151(path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.core_properties.author = "Specs synthetic materials"
    doc.add_heading("SOP-151 Filter Element Change", level=1)
    doc.add_paragraph("Halcyon HX-2 portable fluid pump. Revision B. SYNTHETIC TRAINING DOCUMENT: all "
                      "names, parts, limits and records are fictional demo data.")
    doc.add_heading("1 Scope", level=2)
    doc.add_paragraph("Replace the inline filter element of an HX-2 pump. Use element HX2-420 only.")
    doc.add_heading("2 Warning", level=2)
    doc.add_paragraph("WARNING: Relieve line pressure before opening the filter bowl. Never open the bowl "
                      "while the pump runs.")
    doc.add_heading("3 Torque values", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Fastener", "Size", "Torque")):
        cell.text = text
    for row in (("Bowl clamp screw", "M5", "3 N-m"), ("Drain plug", "M10", "8 N-m")):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_paragraph("Table 1. Torque values for the filter assembly.")
    buf = io.BytesIO()
    filter_figure().save(buf, "PNG")
    buf.seek(0)
    doc.add_picture(buf, width=Inches(4.5))
    doc.add_paragraph("Figure 1. Filter bowl, element and clamp screws.")
    doc.add_heading("4 Procedure", level=2)
    steps = [
        "Stop the pump and close inlet valve V1 and outlet valve V2.",
        "Place a drip tray under the bowl and remove the drain plug to drain about 0.5 L of fluid.",
        "Loosen the two M5 bowl clamp screws and remove the filter bowl.",
        "Replace the filter element with a new element HX2-420 (25 µm rating).",
        "Refit the bowl and torque the two M5 bowl clamp screws to 3 N-m (Table 1).",
        "Refit the drain plug and torque it to 8 N-m (Table 1).",
        "Open V1, then V2, and watch the bowl seam for leaks for 60 s.",
    ]
    for i, s in enumerate(steps, start=1):
        doc.add_paragraph(f"{i}. {s}")
    doc.save(str(path))


# --- SOP-152: phone photo of a printed page ---------------------------------------------

def build_sop152() -> Image.Image:
    page = Image.new("L", (1240, 1754), 250)  # A4 at 150 dpi, slightly off-white paper
    d = ImageDraw.Draw(page)
    y = 110
    def line(text: str, size: int = 30, gap: int = 14) -> None:
        nonlocal y
        d.text((110, y), text, font=font(size), fill=25)
        y += size + gap
    line("SOP-152 Valve Lineup for Flushing", 46, 20)
    line("Halcyon HX-2 portable fluid pump. Revision A. SYNTHETIC TRAINING DOCUMENT.", 24)
    line("All names, parts and limits are fictional demo data.", 24, 40)
    line("WARNING: Wear splash goggles. Never open V2 before V4 is set.", 30, 40)
    for s in ["1. Close suction valve V1 (handle across the pipe).",
              "2. Open bypass valve V3 (handle along the pipe).",
              "3. Set flush valve V4 to 45°.",
              "4. Confirm gauge PG-1 reads below 0.5 bar.",
              "5. Open discharge valve V2 slowly over 10 s."]:
        line(s, 30, 22)
    y += 30
    fig = valve_figure().resize((980, 364))
    page.paste(fig, (130, y))
    y += 380
    line("Figure 1. Valve handle positions for flushing.", 26)
    # Phone photo: slight rotation, uneven light from the top-left, blur and sensor noise.
    rnd = random.Random(152)
    photo = page.rotate(-2.2, resample=Image.BICUBIC, expand=True, fillcolor=70)
    w, h = photo.size
    # Smooth light falloff: bright top-left corner, phone-shadow bottom-right.
    corners = Image.new("L", (2, 2))
    corners.putdata([255, 232, 226, 168])
    shade = corners.resize((w, h), Image.BILINEAR)
    photo = ImageChops.multiply(photo, shade)
    photo = photo.filter(ImageFilter.GaussianBlur(0.8))
    noise = Image.effect_noise((w, h), 12).point(lambda v: v - 128 + 128)
    photo = Image.blend(photo, noise, 0.06)
    return photo.convert("RGB").crop((rnd.randint(10, 30), 20, w - 20, h - 15))


# --- SOP-153: markdown -----------------------------------------------------------------

SOP153 = """---
doc_id: SOP-153
title: Cover Plate Quick Start
revision: A
status: draft
synthetic: true
---
# Halcyon HX-2 portable fluid pump: Cover Plate Quick Start

This fictional training sheet refits the HX-2 inspection cover plate. All names,
parts and limits are synthetic demo data.

> WARNING: Do not overtighten the cover screws; the plastic cover cracks.

## Steps
1. Place the cover plate on the housing with the arrow pointing to the outlet.
2. Start all six M5 cover screws by hand.
3. Tighten the six M5 cover screws to 3 N-m in a star pattern: top (1), bottom (2),
   upper right (3), lower left (4), lower right (5), upper left (6).
4. Check the gap under the cover plate edge with a 0.10 mm feeler gauge. The gauge must not enter.
"""


def main() -> None:
    FIG.mkdir(exist_ok=True)
    seal_housing_figure().save(FIG / "SOP-150-figure-1.png")
    filter_figure().save(FIG / "SOP-151-figure-1.png")
    valve_figure().save(FIG / "SOP-152-figure-1.png")
    (HERE / "SOP-150-pump-seal-swap.pdf").write_bytes(build_sop150())
    build_sop151(HERE / "SOP-151-filter-change.docx")
    build_sop152().save(HERE / "SOP-152-valve-lineup-photo.jpg", quality=85)
    (HERE / "SOP-153-quick-start.md").write_text(SOP153, encoding="utf-8", newline="\n")
    for p in sorted(HERE.glob("SOP-15*")):
        print(f"{p.name:34} {p.stat().st_size:8} bytes")


if __name__ == "__main__":
    main()
